# SPDX-License-Identifier: MIT
# SPDX-FileCopyrightText: 2026 sean galloway
#
# Module: bin/kmaps/styles.py
# Purpose: openpyxl cell styles shared by every signal-contract workbook
"""Cell styles. Kept in one place so every component's workbook looks alike --
a reader moving between stream, pumice and rapids should not have to relearn
what a green cell means."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

GREEN = PatternFill("solid", fgColor="C6EFCE")
GREY = PatternFill("solid", fgColor="F2F2F2")
TITLE = Font(bold=True, size=12)
HDR = Font(bold=True)
MONO = Font(name="Consolas", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Border(*[Side(style="thin")] * 4)

DCFILL = PatternFill("solid", fgColor="FFF2CC")  # don't-care cells
GRAY2 = [(0, 0), (0, 1), (1, 1), (1, 0)]  # gray-code order of 2 bits
