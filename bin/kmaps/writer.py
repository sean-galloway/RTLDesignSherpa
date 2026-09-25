# SPDX-License-Identifier: MIT
# SPDX-FileCopyrightText: 2026 sean galloway
#
# Module: bin/kmaps/writer.py
# Purpose: The contract-table / K-map writer (the three-part required form)
"""KmapWriter and the sheet helpers.

The required artifact is a CONTRACT TABLE, not a grid: term list -> invariants
-> decision table (Sean, 2026-08-28). `kmap()` carries all four evidence
clauses -- axis equations as (name, expr, cite) triples, `depends_only_on`,
`relations=` with a CHECKED predicate, and `rtl_sop` for the derived-vs-RTL
verdict.

The invariant check is this repo's addition over the pumice original: a
predicate that excludes nothing is a claim doing no work, one that excludes
everything is inverted, and both fail the run. Hardware reachability is not
decidable from the map -- what is checked is that the claim bites.

Promoted from the stream generator by TOOLING-KMAP step 5; bodies unchanged.
"""

from .minimize import norm_sop, qm_minimize, sop_str
from .styles import (CENTER, DCFILL, GRAY2, GREEN, GREY, HDR, MONO, THIN,
                     TITLE, WRAP)
from openpyxl.styles import Font

def _glabel(bits):
    return "".join(str(b) for b in bits)


class KmapWriter:
    def __init__(self, ws):
        self.ws = ws
        self.row = 1

    def sheet_intro(self, title, lines):
        ws = self.ws
        ws.cell(self.row, 1, title).font = TITLE
        self.row += 1
        for ln in lines:
            c = ws.cell(self.row, 1, ln)
            c.alignment = WRAP
            self.row += 1
        self.row += 1

    def kmap(self, name, source, expr, varnames, fn, check, values=None,
             depends_only_on=None, rtl_sop=None, relations=None):
        """One K-map block.

        varnames: MSB-first list (2..6). Each entry is either a bare NAME
            (legacy) or a triple (name, defining_expr, "file:line"). Prefer the
            triple: an axis that is itself a composite expression hides exactly
            the logic the map claims to expose, and a bare string cannot tell
            the reader whether the axis is a flop, a wire, or something the
            author invented for the map.

        fn(*bits) -> 0/1, a short string (with `values`), or None meaning
            DON'T-CARE. Unreachable combinations must return None, not a
            convenient 0: a forced 0 both hides a bug (the RTL might drive 1
            there) and blocks a legal simplification (a don't-care is free to
            join an implicant, a 0 is not).

        depends_only_on: why the mapped function ignores every OTHER input.
            A 4-variable map of a 7-variable decision is a SLICE, and a slice
            means nothing without saying what is held constant and why.

        rtl_sop: the RTL expression written as a sum-of-products, so the emitter
            can diff it against the derived minimal cover and render a verdict.

        relations: the SUFFICIENCY half, as (text, reachable_predicate, cite).
            A cell whose bits fail ANY predicate cannot occur in hardware, so it
            is emitted as X (don't-care) instead of a real 0/1 -- a 0 there
            claims the logic was checked in a state it can never reach, and it
            also blocks a legal simplification, since a don't-care is free to
            join an implicant and a 0 is not. predicate=None is an INDEPENDENCE
            note: the pair was examined and is genuinely unconstrained. Saying
            so is part of the argument; silence is not.

            Each constraining relation is CHECKED (TOOLING-KMAP item 0): a
            predicate that excludes nothing is vacuous, and one that excludes
            the whole space is inverted -- both fail the run. Full hardware
            reachability is not decidable from the map, so this checks the
            claim does real work, not that it is true.
        """
        ws = self.ws
        # accept both bare names and (name, expr, cite) triples
        axis_rows = [v for v in varnames if isinstance(v, (tuple, list))]
        varnames = [v[0] if isinstance(v, (tuple, list)) else v for v in varnames]
        ws.cell(self.row, 1, name).font = TITLE
        ws.cell(self.row, 4, source).font = MONO
        self.row += 1
        c = ws.cell(self.row, 1, expr)
        c.font = MONO
        c.alignment = WRAP
        self.row += 1
        n = len(varnames)
        rowv = varnames[0:min(2, n)]                  # grid row variables
        colv = varnames[min(2, n):min(4, n)]          # grid col variables
        pagev = varnames[4:]                          # page variables
        ws.cell(self.row, 1,
                f"rows = {'/'.join(rowv)}   cols = {'/'.join(colv)}"
                + (f"   pages = {'/'.join(pagev)}" if pagev else ""))
        self.row += 1
        ws.cell(self.row, 1, f"CHECK BY INSPECTION: {check}").alignment = WRAP
        ws.cell(self.row, 1).font = Font(italic=True)
        self.row += 1

        # --- axis derivation: what each axis IS, and where it comes from ----
        if axis_rows:
            ws.cell(self.row, 1, "AXES").font = HDR
            self.row += 1
            for hdr, col in (("axis", 1), ("defining expression", 2), ("cite", 4)):
                c = ws.cell(self.row, col, hdr)
                c.font = HDR
            self.row += 1
            for a in axis_rows:
                nm, ex, cite = (list(a) + ["", ""])[:3]
                ws.cell(self.row, 1, nm).font = MONO
                c = ws.cell(self.row, 2, ex)
                c.font = MONO
                c.alignment = WRAP
                ws.cell(self.row, 4, cite).font = MONO
                self.row += 1
        else:
            c = ws.cell(self.row, 1,
                        "AXES: not derived -- axis equations and citations "
                        "missing (see STREAM TASK-001)")
            c.font = Font(italic=True, color="9C6500")
            self.row += 1

        rels = relations or []
        constraining = [r for r in rels if r[1] is not None]
        if rels:
            c = ws.cell(self.row, 1,
                        "RELATIONS between axes (these make cells UNREACHABLE "
                        "-- shown as X, a don't-care, never as 0):")
            c.font = HDR
            c.alignment = WRAP
            self.row += 1
            for text, _pred, cite in rels:
                mark = "    " if _pred is not None else "    (independent) "
                c = ws.cell(self.row, 1, mark + text)
                c.alignment = WRAP
                ws.cell(self.row, 4, cite).font = MONO
                self.row += 1

        def _reachable(bits):
            return all(pr(*bits) for _t, pr, _c in constraining)

        # ---- invariant CHECK (TOOLING-KMAP item 0) ------------------------
        # A relation that excludes nothing is a claim doing no work; one that
        # excludes everything is inverted. Either way the map silently stops
        # meaning what it says, so fail the run rather than emit it. Hardware
        # reachability itself is not decidable here -- this checks the claim
        # bites, not that it is true.
        if constraining:
            _space = [tuple((i >> (n - 1 - k)) & 1 for k in range(n))
                      for i in range(1 << n)]
            for _t, _pr, _c in constraining:
                _excl = [b for b in _space if not _pr(*b)]
                if not _excl:
                    raise SystemExit(
                        f"INVARIANT EXCLUDES NOTHING in kmap {name!r}: "
                        f"{_t!r} ({_c}). Every cell satisfies it, so it is "
                        f"not constraining the map -- drop it or fix the "
                        f"predicate.")
                if len(_excl) == len(_space):
                    raise SystemExit(
                        f"INVARIANT EXCLUDES EVERYTHING in kmap {name!r}: "
                        f"{_t!r} ({_c}). The predicate is inverted -- it must "
                        f"return True where the state IS reachable.")

        if depends_only_on:
            c = ws.cell(self.row, 1, f"DEPENDS ONLY ON: {depends_only_on}")
            c.alignment = WRAP
        else:
            c = ws.cell(self.row, 1,
                        "DEPENDS ONLY ON: not stated -- this map is a SLICE "
                        "with no sufficiency argument (see STREAM TASK-001)")
            c.font = Font(italic=True, color="9C6500")
        self.row += 2

        rows = GRAY2 if len(rowv) == 2 else ([(0,), (1,)] if len(rowv) == 1
                                             else [()])
        cols = GRAY2 if len(colv) == 2 else ([(0,), (1,)] if len(colv) == 1
                                             else [()])
        pages = [()]
        for _ in pagev:
            pages = [p + (b,) for p in pages for b in (0, 1)]

        ones, dontcares = [], []       # minterm indices, for the minimal cover

        for page in pages:
            base = self.row
            if pagev:
                lbl = ", ".join(f"{v}={b}" for v, b in zip(pagev, page))
                ws.cell(base, 1, f"[{lbl}]").font = HDR
                base += 1
            # column headers
            for j, cb in enumerate(cols):
                cc = ws.cell(base, 2 + j, _glabel(cb) or "-")
                cc.font = HDR
                cc.alignment = CENTER
            for i, rb in enumerate(rows):
                rc = ws.cell(base + 1 + i, 1, _glabel(rb) or "-")
                rc.font = HDR
                rc.alignment = CENTER
                for j, cb in enumerate(cols):
                    bits = tuple(rb) + tuple(cb) + tuple(page)
                    v = fn(*bits) if _reachable(bits) else None
                    # minterm index: axis order is rows, cols, pages -- the same
                    # MSB-first order as varnames, so the index matches _qm.
                    idx = 0
                    for b in bits:
                        idx = (idx << 1) | int(bool(b))
                    if v is None:
                        dontcares.append(idx)
                    elif values is None and bool(v):
                        ones.append(idx)
                    cell = ws.cell(base + 1 + i, 2 + j)
                    if v is None:                     # DON'T-CARE
                        cell.value = "X"
                        cell.fill = DCFILL
                    elif values is not None:          # multi-valued map
                        cell.value = values.get(v, str(v))
                        benign = str(v) in ("0", "-", "wait", "hold", "IDLE",
                                            "n/a")
                        cell.fill = GREY if benign else GREEN
                    else:
                        cell.value = int(bool(v))
                        cell.fill = GREEN if v else GREY
                    cell.alignment = CENTER
                    cell.border = THIN
            self.row = base + 1 + len(rows) + 1

        # --- derived minimal cover + verdict vs the RTL --------------------
        # Skipped for multi-valued maps: a minimal SOP is only defined for a
        # boolean function, and forcing one would be a false claim.
        if values is None:
            cubes = qm_minimize(n, ones, dontcares)
            derived = sop_str(cubes, varnames)
            c = ws.cell(self.row, 1, f"DERIVED MINIMAL SOP: {derived}")
            c.font = MONO
            c.alignment = WRAP
            self.row += 1
            if dontcares:
                ws.cell(self.row, 1,
                        f"  ({len(ones)} ones, {len(dontcares)} don't-cares -- "
                        f"X cells were free to widen the cover)").font = Font(
                            italic=True, size=9)
                self.row += 1
            if rtl_sop:
                same = norm_sop(rtl_sop) == norm_sop(derived)
                verdict = ("IDENTICAL -- the RTL is already minimal"
                           if same else
                           "DIFFERS -- reconcile: either the RTL carries "
                           "redundant terms (say why: timing? readability?) "
                           "or an unstated invariant is doing work, or it is "
                           "a bug")
                c = ws.cell(self.row, 1, f"RTL AS WRITTEN:      {rtl_sop}")
                c.font = MONO
                self.row += 1
                c = ws.cell(self.row, 1, f"VERDICT: {verdict}")
                c.font = Font(bold=True, color=("006100" if same else "9C0006"))
                c.alignment = WRAP
            else:
                c = ws.cell(self.row, 1,
                            "VERDICT: NOT CHECKED -- supply rtl_sop= to diff "
                            "the RTL against the derived cover")
                c.font = Font(italic=True, color="9C6500")
            self.row += 1
        self.row += 1

    def table(self, name, source, headers, rows, note=""):
        ws = self.ws
        ws.cell(self.row, 1, name).font = TITLE
        ws.cell(self.row, 4, source).font = MONO
        self.row += 1
        if note:
            c = ws.cell(self.row, 1, note)
            c.alignment = WRAP
            c.font = Font(italic=True)
            self.row += 1
        for j, h in enumerate(headers):
            c = ws.cell(self.row, 1 + j, h)
            c.font = HDR
            c.border = THIN
        self.row += 1
        for r in rows:
            for j, v in enumerate(r):
                c = ws.cell(self.row, 1 + j, v)
                c.border = THIN
                c.alignment = WRAP
            self.row += 1
        self.row += 2


def new_kmap_sheet(wb, name, widths=(30, 9, 9, 9, 9, 9, 9, 9)):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    for col, w in zip("ABCDEFGH", widths):
        ws.column_dimensions[col].width = w
    return KmapWriter(ws)


# ---------------------------------------------------------------------------
# Contract sheets
# ---------------------------------------------------------------------------
CONTRACT_HDRS = ["Group / Channel", "Signal", "Width", "Dir", "Driver",
                 "Legal / Correct behavior (contract)",
                 "Key invariant - assertable check", "Notes / bug history"]


def contract_sheet(wb, name, title, intro, rows):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    for col, w in zip("ABCDEFGH", (16, 28, 10, 6, 16, 58, 46, 46)):
        ws.column_dimensions[col].width = w
    ws.cell(1, 1, title).font = TITLE
    c = ws.cell(2, 1, intro)
    c.alignment = WRAP
    for j, h in enumerate(CONTRACT_HDRS):
        cc = ws.cell(4, 1 + j, h)
        cc.font = HDR
        cc.border = THIN
    r = 5
    for row in rows:
        for j, v in enumerate(row):
            cc = ws.cell(r, 1 + j, v)
            cc.alignment = WRAP
            cc.border = THIN
        r += 1
    return ws


